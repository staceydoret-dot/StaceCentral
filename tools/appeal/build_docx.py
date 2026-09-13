# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = "/tmp/claude-0/-home-user-StaceCentral/a251d0f1-4c8a-560d-85c5-4a48b59404ea/scratchpad/appeal/Stacey-Doret-Appeal-Corrective-Action.docx"
FONT = "Times New Roman"

doc = Document()

# page + base style
for s in doc.sections:
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.element.rPr.rFonts.set(qn("w:cs"), FONT)
pf = normal.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
pf.space_after = Pt(10)


def para(runs, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=10, space_before=0):
    """runs = list of (text, bold, italic)"""
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    for text, b, i in runs:
        r = p.add_run(text)
        r.bold, r.italic = b, i
        r.font.name = FONT
        r.font.size = Pt(12)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        r._element.rPr.rFonts.set(qn("w:cs"), FONT)
    return p


def heading(text):
    return para([(text, True, False)], align=WD_ALIGN_PARAGRAPH.LEFT,
                space_after=6, space_before=12)


L = WD_ALIGN_PARAGRAPH.LEFT

para([("To:\t", True, False),
      ("Human Resources — Employee Relations, Cleveland Clinic", False, False)],
     align=L, space_after=0)
para([("From:\t", True, False), ("Stacey Doret, Employee ID 1026617", False, False)],
     align=L, space_after=0)
para([("Date:\t", True, False), ("September 14, 2026", False, False)],
     align=L, space_after=6)

# horizontal rule
hr = doc.add_paragraph()
hr.paragraph_format.space_after = Pt(6)
from docx.oxml import OxmlElement
pPr = hr._p.get_or_add_pPr()
pbdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "6")
bottom.set(qn("w:space"), "1")
bottom.set(qn("w:color"), "000000")
pbdr.append(bottom)
pPr.append(pbdr)

para([("Re:  Appeal of Corrective Action Issued September 8, 2026", True, False)],
     align=L, space_after=12)

para([("I am submitting this written statement to appeal the corrective action issued to me "
       "on September 8, 2026, within the seven calendar day period provided in the Corrective "
       "Action Employee Acknowledgement.", False, False)])

heading("1.  The Cited Period Falls Within an Approved FMLA Leave Period")
para([("On September 9, 2026, the Absence Management Office issued my Notice of Eligibility "
       "and Designation Notice (Case ID 1000072376) approving intermittent FMLA leave ", False, False),
      ("from August 26, 2026 through on or about August 25, 2027", True, False),
      (". The Designation Notice states: ", False, False),
      ("“APPROVED For FMLA leave. All leave for this reason will be designated as FMLA leave.”", False, True)])
para([("The corrective action was issued September 8, 2026 — ", False, False),
      ("within that approved leave period", True, False),
      (". The Notice of Eligibility further confirms that I notified the Absence Management "
       "Office on September 2, 2026, six days before the corrective action was issued.", False, False)])
para([("I therefore request that any occurrence cited in this corrective action that falls on "
       "or after August 26, 2026 be designated as FMLA leave and removed from my attendance "
       "record.", False, False)])

heading("2.  The Certified Frequency and Duration Covers These Absences")
para([("My provider certified a frequency and duration of ", False, False),
      ("“2 time per week lasting 9 hours per episode,”", False, True),
      (" with appointments at a rate of twice per month. The occurrences at issue are "
       "partial-day late arrivals, which fall well within that certification.", False, False)])
para([("Per the Absence Management Office’s own instructions, absences in closed pay "
       "periods are corrected through a timecard correction request in Workday, and I am "
       "submitting those corrections.", False, False)])

heading("3.  One Occurrence Resulted From a Police Investigation on My Commuter Rail Line")
para([("On Thursday, September 3, 2026, at approximately 6:15 a.m., a fatality occurred on the "
       "Tri-Rail tracks in the 6200 block of North Andrews Avenue in Fort Lauderdale — in "
       "the immediate vicinity of the Cypress Creek station on the line I ride to work — "
       "prompting a Fort Lauderdale Police response and investigation. The incident was reported "
       "by WSVN 7News and Local 10 News that day. My commute that morning was delayed as a "
       "result. This was a public emergency on my transit line, outside my control.", False, False)])

heading("4.  A Reasonable Accommodation Request Was Pending Before This Action Issued")
para([("I initiated a request for a reasonable accommodation regarding my start time on "
       "approximately September 3, 2026 and followed up in writing with Human Resources on "
       "September 4, 2026 — before this corrective action was issued. My manager has since "
       "adjusted my start time by fifteen minutes, confirming that an adjustment is operationally "
       "feasible. My documentation appointment with my treating provider is scheduled for "
       "September 14, 2026.", False, False)])

heading("Requested Outcome")
para([("I ask that the corrective action be rescinded, and specifically that ", False, False),
      ("the associated ineligibility period ending December 8, 2026 be lifted", True, False),
      (", so that I am not barred from applying to internal positions while this matter is "
       "resolved.", False, False)])
para([("Copies of my FMLA Notice of Eligibility and Designation Notice are enclosed. If this "
       "statement should be directed to another representative within Human Resources, please "
       "forward it accordingly and confirm receipt.", False, False)])

para([("Respectfully,", False, False)], align=L, space_after=0, space_before=12)
para([("", False, False)], align=L, space_after=0)
para([("Stacey Doret", False, False)], align=L, space_after=0)
para([("Employee ID 1026617", False, False)], align=L, space_after=12)

para([("Enclosures:  ", True, False),
      ("FMLA Notice of Eligibility (9/9/2026); FMLA Designation Notice (9/9/2026)", False, False)],
     align=L, space_after=0)

doc.save(OUT)
print("WROTE", OUT)
